#!/usr/bin/env python3
"""Gera o documento principal DOCX e PDF. Fontes/tabelas apontam explicitamente limitações."""
from pathlib import Path
import json,csv,subprocess
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from parametros import ROOT,FINANCIALS,ACCESS_DATE
from dossie_base import SECTIONS,DOSSIER_TITLE
from dossie_pts_a import PROGRAMAS
from dossie_pts_b import LIMITATIONS
OUT=ROOT
NAVY='18324B';TEAL='147D78';LIGHT='EAF1F5';AMBER='FFF2CC';GRAY='687785'
def shade(cell,color):
 tcPr=cell._tc.get_or_add_tcPr();shd=OxmlElement('w:shd');shd.set(qn('w:fill'),color);tcPr.append(shd)
def field(paragraph,instr):
 r=OxmlElement('w:fldSimple');r.set(qn('w:instr'),instr);paragraph._p.append(r)
def setup(doc):
 sec=doc.sections[0];sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.68);sec.left_margin=Inches(.78);sec.right_margin=Inches(.78)
 for sty,size,color in [('Normal',9,'2D3943'),('Title',29,NAVY),('Heading 1',17,NAVY),('Heading 2',12,TEAL),('Heading 3',10,NAVY)]:
  s=doc.styles[sty];s.font.name='Aptos' if sty=='Normal' else 'Aptos Display';s.font.size=Pt(size);s.font.color.rgb=RGBColor.from_string(color)
  if sty.startswith('Heading'):s.font.bold=True
  if sty=='Normal':s.paragraph_format.space_after=Pt(6);s.paragraph_format.line_spacing=1.12
 h=sec.header.paragraphs[0];h.text='SIMPAR S.A.  |  DOSSIÊ ACADÊMICO DE AUDITORIA 2025';h.style='Caption';h.runs[0].font.color.rgb=RGBColor.from_string(GRAY)
 f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER;f.add_run('Matrícula 241003357  |  Documento acadêmico — não é opinião de auditoria  |  ');field(f,'PAGE')
def para(doc,text,style=None):
 p=doc.add_paragraph(style=style);p.add_run(text);return p
def table(doc,headers,rows,widths=None,fontsize=8):
 t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for i,h in enumerate(headers):
  c=t.rows[0].cells[i];c.text=str(h);shade(c,NAVY)
  for r in c.paragraphs[0].runs:r.font.bold=True;r.font.color.rgb=RGBColor(255,255,255);r.font.size=Pt(fontsize)
 for row in rows:
  cells=t.add_row().cells
  for i,x in enumerate(row):
   cells[i].text=str(x);cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   for p in cells[i].paragraphs:
    p.paragraph_format.space_after=Pt(2)
    for r in p.runs:r.font.size=Pt(fontsize)
 if widths:
  for row in t.rows:
   for i,w in enumerate(widths):row.cells[i].width=Inches(w)
 return t
def heading(doc,n,text):doc.add_heading(f'{n}. {text}',1)
def add_image(doc,path,caption,width=6.6):
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(path),width=Inches(width));c=doc.add_paragraph(caption,style='Caption');c.alignment=WD_ALIGN_PARAGRAPH.CENTER

doc=Document();setup(doc)
# 1. Cover
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(95);r=p.add_run('SIMPAR S.A.');r.bold=True;r.font.size=Pt(16);r.font.color.rgb=RGBColor.from_string(TEAL)
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run('DOSSIÊ ACADÊMICO\nDE AUDITORIA');r.bold=True;r.font.name='Aptos Display';r.font.size=Pt(32);r.font.color.rgb=RGBColor.from_string(NAVY)
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run('Demonstrações financeiras consolidadas\nExercício encerrado em 31 de dezembro de 2025\nComparativo: 2024').font.size=Pt(15)
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(55);p.add_run('Aluno: Matheus Siqueira\nMatrícula: 241003357\nElaboração: 25 de outubro de 2026\nPortuguês brasileiro').font.size=Pt(12)
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(35);r=p.add_run('TRABALHO ACADÊMICO — NÃO É RELATÓRIO DE AUDITORIA INDEPENDENTE');r.bold=True;r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string('A34B42')
doc.add_page_break()
# 2. Identification
heading(doc,2,'Identificação do trabalho')
table(doc,['Campo','Descrição'],[
('Companhia','SIMPAR S.A.'),('Matrícula','241003357'),('Período','Exercício encerrado em 31/12/2025; comparativo 2024'),('Finalidade','Trabalho acadêmico de auditoria e análise de demonstrações consolidadas.'),('Referencial','Estrutura Conceitual e as 26 NBC TA enumeradas no roteiro (27 critérios no total); versões relevantes ao período.'),('Escopo profissional','Não é auditoria independente. Não se emite opinião nem asseguração.'),('Preparação','25/10/2026; revisão independente não realizada.')],[1.5,5.8])
para(doc,'Convenção de status: Público verificado = informação lida em fonte pública; Planejado = procedimento desenhado, não executado; Simulado = executado apenas em dados artificiais; Pendente = depende de documento ou informação indisponível.',None)
doc.add_page_break()
# 3. TOC
heading(doc,3,'Sumário')
contents=['4. Resumo executivo','5. Objetivo e escopo','6. Metodologia','7. Perfil da SIMPAR','8. Estrutura do grupo e contexto operacional','9. Demonstrações financeiras e comparativos','10. Análise vertical e horizontal','11. Materialidade','12. Riscos','13. Relação norma por norma','14. Programas de auditoria por ciclo','15. Imobilizado e depreciação','16. Receita e contas a receber','17. Caixa e aplicações','18. Estoques','19. Dívidas, arrendamentos e passivos','20. Estimativas e contingências','21. Partes relacionadas','22. Eventos subsequentes','23. Continuidade operacional','24. Consolidação','25. Base sintética','26. Metodologia dos testes','27. Resultados e exceções','28. Gráficos e análises','29. Cobertura normativa','30. Limitações','31. Conclusão integrada','32. Referências','33. Apêndices e papéis de trabalho']
for x in contents:para(doc,x)
para(doc,'Este sumário lista seções, não números de página. A paginação final é apresentada no rodapé.', 'Caption')
# prepared section content lookup
secmap={h:ps for h,ps in SECTIONS}
# 4 summary
heading(doc,4,'Resumo executivo')
for x in secmap['Resumo executivo']:para(doc,x)
# 5 objective/scope
heading(doc,5,'Objetivo e escopo')
for x in secmap['Objetivo, escopo e metodologia'][:2]:para(doc,x)
# 6 methodology
heading(doc,6,'Metodologia')
for x in secmap['Objetivo, escopo e metodologia'][2:]:para(doc,x)
para(doc,'Fluxo de trabalho proposto. Fontes públicas informam o planejamento, mas não substituem registros internos, confirmações e inspeções que continuam pendentes; a base sintética é somente didática e não substitui evidência sobre a Companhia.',None)
doc.add_page_break()
add_image(doc,ROOT/'graficos'/'fluxo_auditoria.png','Figura 1. Processo de auditoria proposto. Diagrama próprio, renderizado do fonte Mermaid arquivado em graficos/.',6.4)
doc.add_page_break()
# 7/8
heading(doc,7,'Perfil da SIMPAR')
for x in secmap['Perfil do grupo e contexto operacional'][:1]:para(doc,x)
heading(doc,8,'Estrutura do grupo e contexto operacional')
for x in secmap['Perfil do grupo e contexto operacional'][1:]:para(doc,x)
# 9
heading(doc,9,'Demonstrações financeiras de 2025 e comparativo de 2024')
for x in secmap['Demonstrações e análise horizontal'][:3]:para(doc,x)
finrows=[]
for nm in ['Ativo total','Imobilizado líquido','Receita líquida','Lucro antes de IR/CS','Lucro de operações continuadas','Lucro líquido do exercício','Caixa e equivalentes','Fluxo operacional líquido','Fluxo de investimento líquido','Fluxo de financiamento líquido']:
 a,b=FINANCIALS[nm]; delta=a-b; pct=(a/b-1)*100 if b>0 else None
 finrows.append((nm,f'{a:,}'.replace(',','.'),f'{b:,}'.replace(',','.'),f'{delta:+,}'.replace(',','.'),f'{pct:+.2f}%' if pct is not None else 'n/a (base 2024 negativa)'))
table(doc,['Indicador (R$ mil)','2025','2024','Variação absoluta','Variação %'],finrows,[2.35,.9,.9,1.1,1.1],7)
for x in secmap['Demonstrações e análise horizontal'][3:]:para(doc,x)
para(doc,'Relatório do auditor: o resumo publicado no PDF consultado informa opinião “sem modificação”, assinado pela PwC em 30/03/2026 (PDF p. 4; impresso B15). O próprio resumo remete ao relatório completo divulgado eletronicamente. O arquivo integral não foi obtido; portanto, os principais assuntos de auditoria, eventuais ênfases e texto completo não foram avaliados. A consulta RAD também identifica PwC como auditora [1] [2].',None)
para(doc,'Informações por segmento: a estrutura do grupo está descrita no Relatório Anual Integrado, mas as informações financeiras segmentadas e reconciliações das notas da DFP integral não foram analisadas em profundidade nesta cópia resumida. A análise comparativa é consolidada; segmentação é procedimento pendente [2] [4].',None)
# 10
heading(doc,10,'Análise vertical e horizontal')
add_image(doc,ROOT/'graficos'/'comparativo_2024_2025.png','Figura 2. Comparativo consolidado selecionado, valores em R$ bilhões, convertidos de R$ mil; FCO permaneceu negativo. Fonte [2], PDF p. 2.',6.5)
metric=json.loads((ROOT/'dados'/'indicadores.json').read_text())
ratio_rows=[('Imobilizado / ativo total',metric['Imobilizado/Ativo total 2025'],metric['Imobilizado/Ativo total 2024']),('Caixa / ativo total',metric['Caixa/Ativo total 2025'],metric['Caixa/Ativo total 2024']),('Margem bruta',metric['Margem bruta 2025'],metric['Margem bruta 2024']),('Margem antes de IR/CS',metric['Margem lucro antes de IR/CS 2025'],metric['Margem lucro antes de IR/CS 2024']),('Liquidez corrente',metric['Liquidez corrente 2025'],metric['Liquidez corrente 2024'])]
table(doc,['Indicador vertical/razão','2025','2024'],[(n,f'{a:.2%}' if a<1 else f'{a:.2f}x',f'{b:.2%}' if b<1 else f'{b:.2f}x') for n,a,b in ratio_rows],[3.2,1.4,1.4])
para(doc,'Liquidez corrente foi calculada como ativo circulante ÷ passivo circulante: 28.898.090 ÷ 24.251.860 = 1,19x em 2025; 27.220.964 ÷ 21.080.863 = 1,29x em 2024. Indicadores calculados a partir de rubricas agregadas, não índices de covenants. Variações percentuais não são exibidas para comparativos de fluxos negativos com sinal instável.',None)
# 11
heading(doc,11,'Materialidade')
for x in secmap['Materialidade preliminar']:para(doc,x)
mat=json.loads((ROOT/'dados'/'materialidade.json').read_text())
table(doc,['Parâmetro','Fórmula','R$ mil'],[
('Materialidade global','43.528.375 × 0,5%',f'{mat["global_R$_mil"]:,.3f}'),('Execução','Global × 75%',f'{mat["execution_R$_mil"]:,.3f}'),('Claramente trivial','Global × 5%',f'{mat["clearly_trivial_R$_mil"]:,.3f}'),('Alternativa: ativo','88.642.122 × 0,5%',f'{mat["alternative_benchmarks_R$_mil"]["0.5% do ativo total"]:,.3f}'),('Alternativa: lucro antes IR/CS','254.114 × 5%',f'{mat["alternative_benchmarks_R$_mil"]["5% do lucro antes de IR/CS"]:,.3f}'),('Alternativa: PL','7.411.001 × 1%',f'{mat["alternative_benchmarks_R$_mil"]["1% do patrimônio líquido"]:,.3f}')],[2.0,2.4,1.4],7)
add_image(doc,ROOT/'graficos'/'materialidade.png','Figura 3. Materialidade preliminar versus saldos; R$ milhões; percentuais sujeitos a julgamento.',6.5)
# 12 risk
heading(doc,12,'Identificação e avaliação preliminar dos riscos')
para(doc,'A matriz detalhada contém 19 riscos em nove ciclos. Probabilidade e impacto são julgamentos preliminares de planejamento, não resultados de testes.',None)
with (ROOT/'matrizes'/'matriz_riscos.csv').open(encoding='utf-8-sig') as f: risks=list(csv.DictReader(f))
table(doc,['Ciclo','Risco','Afirmações','Probabilidade','Impacto','Status'],[(r['Ciclo'],r['Risco'],r['Afirmações'],r['Probabilidade preliminar'],r['Impacto preliminar'],r['Status']) for r in risks],[1.1,1.7,1.6,.8,.6,.7],6.5)
para(doc,'Cobertura: imobilizado/depreciação; receita/recebíveis; caixa; estoques; dívida/arrendamentos; passivos não registrados; estimativas/contingências; partes relacionadas; consolidação; continuidade. Respostas e evidências constam em matrizes/matriz_riscos.csv.',None)
# 13 matrix 28
heading(doc,13,'Relação norma por norma')
para(doc,'A matriz abaixo relaciona a Estrutura Conceitual e as 26 NBC TA enumeradas no roteiro (27 critérios no total) à SIMPAR. A versão aplicável é indicada separadamente da versão publicada após o período; 240(R2) e 570(R1) são revisões de 2026 com aplicação prospectiva. Papéis individuais PT-EC a PT-600 acompanham o ZIP.',None)
with (ROOT/'matrizes'/'matriz_cobertura_normas.csv').open(encoding='utf-8-sig') as f: norms=list(csv.DictReader(f))
norm_rows=[]
for r in norms:
 brief=r['Aplicação SIMPAR 2025'].split(';')[0].strip()
 if len(brief)>105:brief=brief[:102].rsplit(' ',1)[0]+'…'
 norm_rows.append((r['Código'],r['Norma/revisão']+'\n'+r['Versão aplicável ao período'],brief+'\nStatus: '+r['Status da evidência']))
table(doc,['Código','Norma/revisão e versão do período','Aplicação resumida e status'],norm_rows,[.55,2.15,4.0],7.0)
para(doc,'Revisar no CFC os textos integrais e itens específicos antes de uso acadêmico definitivo. A descrição é síntese de planejamento; não transcrição literal. Fontes oficiais individuais estão no CSV de cobertura e nos papéis.',None)
# 14 programs
heading(doc,14,'Programas de auditoria por ciclo')
for name,items in PROGRAMAS.items():
 doc.add_heading(name,2)
 for h,b in items:
  p=doc.add_paragraph();p.add_run(h+': ').bold=True;p.add_run(b)
# 15-24 detailed topics
headings={15:'Imobilizado e depreciação',16:'Receita e contas a receber',17:'Caixa e aplicações financeiras',18:'Estoques',19:'Dívidas, arrendamentos e passivos',20:'Estimativas, provisões e contingências',21:'Partes relacionadas',22:'Eventos subsequentes',23:'Continuidade operacional',24:'Processo de consolidação do grupo'}
for n,h in headings.items():
 heading(doc,n,h)
 key={15:'Imobilizado, depreciação e estimativas',16:'Receita, contas a receber, estoques e caixa',17:'Receita, contas a receber, estoques e caixa',18:'Receita, contas a receber, estoques e caixa',19:'Dívida, arrendamentos, contingências e partes relacionadas',20:'Dívida, arrendamentos, contingências e partes relacionadas',21:'Dívida, arrendamentos, contingências e partes relacionadas',22:'Consolidação e continuidade operacional',23:'Consolidação e continuidade operacional',24:'Consolidação e continuidade operacional'}[n]
 vals=secmap[key]
 if n==15:sel=vals
 elif n==16:sel=secmap['Receita, contas a receber, estoques e caixa'][:2]
 elif n==17:sel=secmap['Receita, contas a receber, estoques e caixa'][2:]
 elif n==18:sel=secmap['Receita, contas a receber, estoques e caixa'][1:2]
 elif n==19:sel=secmap['Dívida, arrendamentos, contingências e partes relacionadas'][:1]
 elif n==20:sel=secmap['Dívida, arrendamentos, contingências e partes relacionadas'][1:2]
 elif n==21:sel=secmap['Dívida, arrendamentos, contingências e partes relacionadas'][2:]
 elif n==22:sel=secmap['Consolidação e continuidade operacional'][2:]
 elif n==23:sel=secmap['Consolidação e continuidade operacional'][1:2]
 else:sel=secmap['Consolidação e continuidade operacional'][:1]
 for x in sel:para(doc,x)
 # Include a clear procedure-response reference
 if n in [15,16,17,18,19,20,21,24]:para(doc,'Resposta proposta: consultar PT-04/matriz_riscos.csv e programa por ciclo na seção 14. Todos os testes dependentes de registros internos estão PLANEJADOS ou PENDENTES.',None)
# 25-27 synthetic
heading(doc,25,'Base sintética de 500 mil registros')
for x in secmap['Base sintética e testes de dados'][:2]:para(doc,x)
heading(doc,26,'Metodologia dos testes de dados')
for x in secmap['Base sintética e testes de dados'][2:]:para(doc,x)
para(doc,'Rotina executada: `python3 codigo/01_gerar_base.py` e `python3 codigo/02_testes_base.py`. `codigo/03_analiticos.py` prepara tabelas comparativas e materialidade; `codigo/04_graficos.py` prepara imagens.',None)
heading(doc,27,'Resultados e exceções da base sintética')
for x in (ROOT/'dados'/'resultados_testes.txt').read_text().splitlines():
 if x:para(doc,x)
add_image(doc,ROOT/'graficos'/'distribuicao_sintetica.png','Figura 4. Distribuição sintética por classe; valores artificiais, não SIMPAR.',6.1)
add_image(doc,ROOT/'graficos'/'anomalias_sinteticas.png','Figura 5. Contagens deliberadamente inseridas/detectadas; não são exceções reais.',6.1)
# 28
heading(doc,28,'Gráficos e análises')
para(doc,'Os gráficos financeiros usam somente cifras publicadas em R$ mil e conversão para R$ bilhões, com fonte no respectivo rodapé. Os gráficos sintéticos identificam explicitamente a artificialidade. Nenhum dado faltante foi extrapolado para preencher séries.',None)
# 29
heading(doc,29,'Matriz de cobertura das NBC TA')
para(doc,f'Cobertura: {len(norms)} critérios (Estrutura Conceitual + 26 NBC TA); arquivo editável e URL por critério em `matrizes/matriz_cobertura_normas.csv`. A matriz demonstra planejamento e status de evidência, não conformidade completa do trabalho.',None)
# 30
heading(doc,30,'Limitações e procedimentos pendentes')
for code,area,lim,status,action in LIMITATIONS:
 p=doc.add_paragraph();p.add_run(f'{code} — {area} ({status}): ').bold=True;p.add_run(lim+' Ação: '+action)
# 31
heading(doc,31,'Conclusão integrada')
for x in secmap['Conclusão integrada e limitações']:para(doc,x)
para(doc,'Conclusão: evidência pública suficiente para contextualizar valores publicados; insuficiente para opinião sobre demonstrações ou conclusão sobre registros individuais, controles ou auditoria de grupo. Não se emite opinião de auditor independente.',None)
# 32 references
heading(doc,32,'Referências')
refs=[
'[1] Comissão de Valores Mobiliários. SIMPAR S.A., DFP 31/12/2025, V1, documento sequencial 155854. https://www.rad.cvm.gov.br/ENET/frmGerenciaPaginaFRE.aspx?CodigoTipoInstituicao=1&NumeroSequencialDocumento=155854 (acesso 25/10/2026). Metadados consultados; tabelas/PDF integral indisponíveis na interface utilizada.',
'[2] SIMPAR S.A. Demonstrações financeiras individuais e consolidadas, exercícios findos em 31/12/2025 e 31/12/2024, R$ mil. Cópia reproduzida por O Estado de S. Paulo, publicação 31/03/2026, PDF consultado p. 2; notas 2.12 (política de imobilizado), 2.14 (financiamentos), 2.18 (provisões), 2.20 (partes relacionadas), 6 (eventos subsequentes). https://estadaori.estadao.com.br/wp-content/uploads/2026/03/simpar-sa-balanco-2026-03-31_07-36-21.pdf. Fonte secundária/reprodução; conferir no RAD original.',
'[3] Comissão de Valores Mobiliários. SIMPAR S.A., DFP 31/12/2024, V2, documento sequencial 146257. https://www.rad.cvm.gov.br/ENET/frmGerenciaPaginaFRE.aspx?CodigoTipoInstituicao=1&NumeroSequencialDocumento=146257 (acesso 25/10/2026).',
'[4] SIMPAR S.A. Relatório Anual Integrado 2025, publicado em 31/08/2026; seção resultados econômico-financeiros pp. 90–91; páginas do PDF. https://simpar.com.br/wp-content/uploads/2026/09/Relatorio-Anual-Integrado-2025.pdf (acesso 25/10/2026). Fonte oficial da companhia; não substitui DFs auditadas.',
'[5] Conselho Federal de Contabilidade. NBC TA — Auditoria Independente de Informação Contábil Histórica, listagem oficial. https://cfc.org.br/tecnica/normas-brasileiras-de-contabilidade/nbc-ta-de-auditoria-independente/ (acesso 25/10/2026).',
'[6] Conselho Federal de Contabilidade. NBC TA 240 (R1), Resolução 2016/NBCTA240(R1); e NBC TA 240 (R2), Resolução 2026/NBCTA240(R2), DOU 25/09/2026. https://www2.cfc.org.br/sisweb/sre/detalhes_sre.aspx?Codigo=2016/NBCTA240%28R1%29 ; https://www2.cfc.org.br/sisweb/sre/detalhes_sre.aspx?Codigo=2026/NBCTA240%28R2%29 (acesso 25/10/2026).',
'[7] Conselho Federal de Contabilidade. NBC TA 570, Resolução 2016/NBCTA570; e NBC TA 570 (R1), Resolução 2026/NBCTA570(R1), DOU 25/09/2026. https://www2.cfc.org.br/sisweb/sre/detalhes_sre.aspx?Codigo=2016/NBCTA570 ; https://www2.cfc.org.br/sisweb/sre/detalhes_sre.aspx?Codigo=2026/NBCTA570%28R1%29 (acesso 25/10/2026).',
'[8] Conselho Federal de Contabilidade. NBC TA 600 (R2), Resolução 2023/NBCTA600(R2), DOU 06/12/2023. https://www2.cfc.org.br/sisweb/sre/detalhes_sre.aspx?Codigo=2023/NBCTA600%28R2%29 (acesso 25/10/2026).',
'[9] Dados sintéticos gerados neste pacote: `codigo/01_gerar_base.py`; testes: `codigo/02_testes_base.py`; método/seed documentados em `dados/dicionario_variaveis.md`. Não é fonte da SIMPAR.'
]
for x in refs:para(doc,x)
# 33 appendices
heading(doc,33,'Apêndices e papéis de trabalho')
para(doc,'A. PT-EC e PT-200 a PT-600: um Word por critério normativo. B. PT-01 a PT-05: planejamento, base sintética, programas, riscos e conclusão. C. Matriz de riscos, cobertura normativa, fontes, procedimentos e limitações em CSV/Excel. D. Dados artificiais e códigos reproduzíveis. E. Gráficos e fluxograma próprios. F. PDF consultado das demonstrações e Relatório Anual Integrado preservados em `fontes/`.',None)
para(doc,'Condição de uso: academic review only. Cópias de demonstrações e relatórios pertencem a seus titulares; conservar citações e verificar condições de uso antes de redistribuição.',None)
# Save
path=OUT/'Dossie_Auditoria_SIMPAR_2025.docx';doc.save(path)
# Convert DOCX to PDF via LibreOffice; fail loudly if not possible.
res=subprocess.run(['libreoffice','--headless','--convert-to','pdf','--outdir',str(OUT),str(path)],capture_output=True,text=True,timeout=120)
print(res.stdout,res.stderr)
pdf=OUT/'Dossie_Auditoria_SIMPAR_2025.pdf'
if res.returncode or not pdf.exists():raise SystemExit(f'Falha ao converter para PDF: {res.returncode}')
print('Gerados:',path,pdf)
