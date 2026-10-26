#!/usr/bin/env python3
"""Valida integridade básica dos entregáveis reconstruídos; não valida opinião nem evidenciação de saldos."""
from pathlib import Path
import csv,subprocess
from docx import Document
import openpyxl
from zipfile import ZipFile
R=Path(__file__).resolve().parents[1]
def check(cond,msg):
 if not cond:raise SystemExit('FALHOU: '+msg)
 print('OK:',msg)
# Base sintética e testes
with (R/'dados'/'base_imobilizado_sintetica.csv').open(encoding='utf-8-sig',newline='') as f:n=sum(1 for _ in f)-1
check(n==500_000,f'CSV tem 500.000 registros (atual={n})')
txt=(R/'dados'/'resultados_testes.txt').read_text(encoding='utf-8')
check(txt.count('OK')>=12,'12 verificações da base concluíram OK')
check('inserida=3; detectada=3; OK' in txt and 'inserida=2; detectada=2; OK' in txt,'anomalias deliberadas conferem com as regras de teste')
# Matriz
with (R/'matrizes'/'matriz_cobertura_normas.csv').open(encoding='utf-8-sig') as f:norms=list(csv.DictReader(f))
check(len(norms)==27,'matriz cobre 27 critérios (EC + 26 NBC TA)')
check((R/'matrizes'/'Matriz_Cobertura_Normas.xlsx').exists(),'matriz normativa Excel autônoma foi gerada')
# Workpapers and main docx
wps=list((R/'workpapers').glob('*.docx'))
check(len(wps)==32,f'32 papéis de trabalho DOCX (atual={len(wps)})')
for p in wps:check(bool(Document(p).paragraphs),f'DOCX abre: {p.name}')
main=Document(R/'Dossie_Auditoria_SIMPAR_2025.docx')
check(len(main.inline_shapes)>=5,'DOCX principal inclui gráficos/diagrama')
# XLSX integrity/read
wb=openpyxl.load_workbook(R/'dados'/'Dossie_SIMPAR_Analises.xlsx',read_only=True,data_only=False)
check('Cobertura NBC TA' in wb.sheetnames,'planilha principal contém aba de cobertura normativa')
wb2=openpyxl.load_workbook(R/'matrizes'/'Matriz_Cobertura_Normas.xlsx',read_only=True)
check(wb2['Cobertura NBC TA'].max_row==28,'matriz Excel tem cabeçalho e 27 critérios')
# PDF, zip formats
pdf=R/'Dossie_Auditoria_SIMPAR_2025.pdf';info=subprocess.run(['pdfinfo',str(pdf)],capture_output=True,text=True,check=True).stdout
check('Pages:' in info and pdf.stat().st_size>100_000,'PDF abre e tem conteúdo')
for p in [R/'Dossie_Auditoria_SIMPAR_2025.docx',R/'dados'/'Dossie_SIMPAR_Analises.xlsx',R/'matrizes'/'Matriz_Cobertura_Normas.xlsx']:
 with ZipFile(p) as z:check(z.testzip() is None,f'arquivo compactado íntegro: {p.name}')
print('Validação mecânica concluída; não substitui revisão humana, obtenção do DFP original ou execução de procedimentos de auditoria.')
